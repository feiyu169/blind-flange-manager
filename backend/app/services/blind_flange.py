"""盲板服务"""

import io
from typing import List, Optional

from fastapi import HTTPException, UploadFile, status
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.blind_flange import BlindFlange, BlindFlangeStatusLog
from app.schemas.blind_flange import (
    BlindFlangeCreate,
    BlindFlangeImportItem,
    BlindFlangeImportResponse,
    BlindFlangeUpdate,
)


async def create_blind_flange(
    db: AsyncSession,
    data: BlindFlangeCreate,
    operator_id: int,
) -> BlindFlange:
    """创建盲板"""
    # 检查编码是否已存在
    result = await db.execute(select(BlindFlange).where(BlindFlange.code == data.code))
    if result.scalar_one_or_none():
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"盲板编码 {data.code} 已存在",
        )

    blind_flange = BlindFlange(**data.model_dump())
    db.add(blind_flange)
    await db.flush()

    # 记录状态变更日志
    if data.status:
        log = BlindFlangeStatusLog(
            blind_flange_id=blind_flange.id,
            old_status=None,
            new_status=data.status,
            operator_id=operator_id,
            reason="创建盲板",
        )
        db.add(log)

    await db.flush()
    await db.refresh(blind_flange)
    return blind_flange


async def get_blind_flange(
    db: AsyncSession,
    blind_flange_id: int,
) -> BlindFlange:
    """获取盲板详情"""
    result = await db.execute(select(BlindFlange).where(BlindFlange.id == blind_flange_id))
    blind_flange = result.scalar_one_or_none()
    if not blind_flange:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="盲板不存在",
        )
    return blind_flange


async def get_blind_flange_by_code(
    db: AsyncSession,
    code: str,
) -> BlindFlange:
    """根据编码获取盲板"""
    result = await db.execute(select(BlindFlange).where(BlindFlange.code == code))
    blind_flange = result.scalar_one_or_none()
    if not blind_flange:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"盲板编码 {code} 不存在",
        )
    return blind_flange


async def list_blind_flanges(
    db: AsyncSession,
    page: int = 1,
    page_size: int = 20,
    status: Optional[str] = None,
    keyword: Optional[str] = None,
) -> dict:
    """获取盲板列表"""
    query = select(BlindFlange)

    # 状态筛选
    if status:
        query = query.where(BlindFlange.status == status)

    # 关键词搜索
    if keyword:
        query = query.where(
            BlindFlange.code.ilike(f"%{keyword}%")
            | BlindFlange.name.ilike(f"%{keyword}%")
            | BlindFlange.location.ilike(f"%{keyword}%")
            | BlindFlange.pipeline_no.ilike(f"%{keyword}%")
        )

    # 计算总数
    count_query = select(func.count()).select_from(query.subquery())
    total_result = await db.execute(count_query)
    total = total_result.scalar()

    # 分页查询
    query = query.order_by(BlindFlange.id.desc())
    query = query.offset((page - 1) * page_size).limit(page_size)
    result = await db.execute(query)
    items = result.scalars().all()

    return {
        "items": items,
        "total": total,
        "page": page,
        "page_size": page_size,
    }


async def update_blind_flange(
    db: AsyncSession,
    blind_flange_id: int,
    data: BlindFlangeUpdate,
    operator_id: int,
) -> BlindFlange:
    """更新盲板"""
    result = await db.execute(select(BlindFlange).where(BlindFlange.id == blind_flange_id))
    blind_flange = result.scalar_one_or_none()
    if not blind_flange:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="盲板不存在",
        )

    update_data = data.model_dump(exclude_unset=True)

    # 记录状态变更
    if "status" in update_data and update_data["status"] != blind_flange.status:
        log = BlindFlangeStatusLog(
            blind_flange_id=blind_flange.id,
            old_status=blind_flange.status,
            new_status=update_data["status"],
            operator_id=operator_id,
            reason=update_data.get("notes", "状态变更"),
        )
        db.add(log)

    for key, value in update_data.items():
        setattr(blind_flange, key, value)

    await db.flush()
    await db.refresh(blind_flange)
    return blind_flange


async def delete_blind_flange(
    db: AsyncSession,
    blind_flange_id: int,
) -> None:
    """删除盲板"""
    result = await db.execute(select(BlindFlange).where(BlindFlange.id == blind_flange_id))
    blind_flange = result.scalar_one_or_none()
    if not blind_flange:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="盲板不存在",
        )

    await db.delete(blind_flange)
    await db.flush()


async def get_status_logs(
    db: AsyncSession,
    blind_flange_id: int,
) -> List[BlindFlangeStatusLog]:
    """获取状态变更日志"""
    result = await db.execute(
        select(BlindFlangeStatusLog)
        .where(BlindFlangeStatusLog.blind_flange_id == blind_flange_id)
        .order_by(BlindFlangeStatusLog.created_at.desc())
    )
    return result.scalars().all()


async def import_blind_flanges(
    db: AsyncSession,
    file: UploadFile,
    operator_id: int,
) -> BlindFlangeImportResponse:
    """批量导入盲板"""
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="只支持 Excel 文件",
        )

    content = await file.read()
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active

    success_count = 0
    fail_count = 0
    errors = []

    # 跳过表头
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        try:
            if not row or not row[0]:
                continue

            data = BlindFlangeImportItem(
                code=str(row[0]),
                name=str(row[1]),
                specification=str(row[2]) if len(row) > 2 and row[2] else None,
                material=str(row[3]) if len(row) > 3 and row[3] else None,
                pressure_rating=str(row[4]) if len(row) > 4 and row[4] else None,
                size=str(row[5]) if len(row) > 5 and row[5] else None,
                location=str(row[6]) if len(row) > 6 and row[6] else None,
                pipeline_no=str(row[7]) if len(row) > 7 and row[7] else None,
                flange_no=str(row[8]) if len(row) > 8 and row[8] else None,
                notes=str(row[9]) if len(row) > 9 and row[9] else None,
            )

            # 检查编码是否已存在
            result = await db.execute(select(BlindFlange).where(BlindFlange.code == data.code))
            if result.scalar_one_or_none():
                errors.append(f"第 {row_idx} 行: 编码 {data.code} 已存在")
                fail_count += 1
                continue

            blind_flange = BlindFlange(**data.model_dump())
            db.add(blind_flange)
            success_count += 1

        except Exception as e:
            errors.append(f"第 {row_idx} 行: {str(e)}")
            fail_count += 1

    await db.flush()

    return BlindFlangeImportResponse(
        success_count=success_count,
        fail_count=fail_count,
        errors=errors,
    )


async def export_blind_flanges(
    db: AsyncSession,
    status: Optional[str] = None,
) -> io.BytesIO:
    """导出盲板"""
    query = select(BlindFlange)

    if status:
        query = query.where(BlindFlange.status == status)

    query = query.order_by(BlindFlange.id)
    result = await db.execute(query)
    items = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "盲板列表"

    # 表头
    headers = [
        "编码", "名称", "规格型号", "材质", "压力等级", "尺寸",
        "安装位置", "管道编号", "法兰编号", "状态", "备注",
        "创建时间", "更新时间",
    ]
    ws.append(headers)

    # 数据
    for item in items:
        ws.append([
            item.code,
            item.name,
            item.specification,
            item.material,
            item.pressure_rating,
            item.size,
            item.location,
            item.pipeline_no,
            item.flange_no,
            item.status,
            item.notes,
            item.created_at.strftime("%Y-%m-%d %H:%M:%S") if item.created_at else None,
            item.updated_at.strftime("%Y-%m-%d %H:%M:%S") if item.updated_at else None,
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
